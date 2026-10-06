"""Testes do capítulo 13 (onde colocar fiscal): critérios, níveis, risco e exportáveis.

Fixtures pequenas, sem rede e sem os bancos grandes. Os dois últimos testes leem os
arquivos publicados (`fiscais.json` e o Excel) quando eles existem.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from apuracao_2026 import fiscais_base as fb
from apuracao_2026 import fiscais_criterios as fc
from apuracao_2026 import fiscais_export as fx
from apuracao_2026 import fiscais_regras as fr
from apuracao_2026 import fiscais_territorio as ft
from apuracao_2026 import fiscais_texto as ftx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
JSON = ROOT / "analysis/apuracao_2026/dados/fiscais.json"
XLSX = ROOT / "docs/assets/fiscais_2026.xlsx"
CSV = ROOT / "docs/assets/fiscais_2026.csv"


def secao(**kw: object) -> dict[str, object]:
    base: dict[str, object] = {
        "uf": "sp",
        "mun": "71072",
        "zona": 1,
        "secao": 1,
        "local_id": "SP-71072-1-1000",
        "v13": 100,
        "v22": 100,
        "validos": 200,
        "brancos": 4,
        "nulos": 6,
        "comparecimento": 210,
        "aptos": 260,
        "lula_pct": 50.0,
        "flavio_pct": 50.0,
        "zona_resto_lula_pct": 50.0,
        "zona_resto_flavio_pct": 50.0,
        "uf_lula_pct": 50.0,
        "uf_flavio_pct": 50.0,
        "lula22_pct": np.nan,
        "bolso22_pct": np.nan,
        "lula22_2t_pct": np.nan,
        "bolso22_2t_pct": np.nan,
        "nominais_1t": np.nan,
        "mesma_secao": False,
        "tipo_arquivo": 1,
        "tipo_urna": 1,
        "n_cargas": 1,
        "_enc": pd.Timestamp("2026-10-04 17:05:00"),
        "_dr": pd.Timestamp("2026-10-04 18:30:00"),
        "tipo_inferido": "escola ou universidade",
        "tipo_local": "Convencional",
        "aptos_tte": 0,
    }
    base.update(kw)
    return base


def quadro(*linhas: dict[str, object]) -> pd.DataFrame:
    return pd.DataFrame(list(linhas))


# ---------------------------------------------------------------- critérios


def test_criterio_a_exige_excesso_sobre_a_zona_e_marca_enclave() -> None:
    df = quadro(
        secao(lula_pct=95.0, flavio_pct=3.0, zona_resto_lula_pct=70.0),
        secao(lula_pct=95.0, flavio_pct=3.0, zona_resto_lula_pct=80.0),
        secao(
            lula_pct=95.0, flavio_pct=3.0, zona_resto_lula_pct=60.0, comparecimento=90
        ),
        secao(
            lula_pct=96.0,
            flavio_pct=2.0,
            zona_resto_lula_pct=60.0,
            mesma_secao=True,
            lula22_pct=93.0,
            lula22_2t_pct=95.0,
        ),
        secao(lula_pct=4.0, flavio_pct=92.0, zona_resto_flavio_pct=60.0),
    )
    a = fc.crit_a(df)
    assert a["a"].tolist() == [True, False, False, True, True]
    assert a["enclave_2022"].tolist() == [False, False, False, True, False]
    assert a.loc[4, "a_cand"] == "flavio"
    assert a.loc[0, "a_exc"] == pytest.approx(25.0)


def test_criterios_b_c_g() -> None:
    df = quadro(
        secao(v22=0, comparecimento=200),
        secao(v13=0, comparecimento=120),
        secao(comparecimento=260, aptos=260),
        secao(comparecimento=40, aptos=40),
        secao(_dr=pd.Timestamp("2026-10-05 00:10:00")),
        secao(uf="zz", _dr=pd.Timestamp("2026-10-05 03:00:00")),
    )
    assert fc.crit_b(df)["b"].tolist() == [True, False, False, False, False, False]
    assert fc.crit_b(df).loc[0, "b_quem"] == "flavio"
    assert fc.crit_c(df)["c"].tolist() == [False, False, True, False, False, False]
    assert fc.crit_g(df)["g"].tolist() == [False, False, False, False, True, False]


def test_criterio_d_hora_de_brasilia_e_excesso_de_lula() -> None:
    df = quadro(
        secao(_enc=pd.Timestamp("2026-10-04 19:10:00"), lula_pct=70.0),
        secao(_enc=pd.Timestamp("2026-10-04 19:10:00"), lula_pct=55.0),
        secao(_enc=pd.Timestamp("2026-10-04 18:50:00"), lula_pct=70.0),
        secao(_enc=pd.NaT, lula_pct=70.0),
    )
    assert fc.crit_d(df)["d"].tolist() == [True, False, False, False]


def test_criterio_f_exige_urna_fora_do_padrao_e_diferenca() -> None:
    df = quadro(
        secao(tipo_urna=4, lula_pct=62.0),
        secao(tipo_urna=4, lula_pct=55.0, flavio_pct=45.0),
        secao(n_cargas=2, flavio_pct=38.0),
        secao(lula_pct=70.0),
    )
    assert fc.crit_f(df)["f"].tolist() == [True, False, True, False]


def test_criterio_i_so_nas_zonas_do_topo() -> None:
    topo = {("sp", "71072", 1): {"posicao": 3, "escore": 99.5}}
    df = quadro(
        secao(lula_pct=62.0),
        secao(lula_pct=58.0),
        secao(zona=2, lula_pct=80.0),
    )
    assert fc.crit_i(df, topo)["i"].tolist() == [True, False, False]


def test_criterio_j_tres_desvios_da_uf() -> None:
    rng = np.random.default_rng(1)
    linhas = []
    for k in range(200):
        sw = float(rng.normal(0, 2))
        linhas.append(
            secao(
                secao=k + 1,
                mesma_secao=True,
                nominais_1t=200,
                lula22_pct=50.0,
                bolso22_pct=50.0,
                lula_pct=50.0 - sw / 2,
                flavio_pct=50.0 + sw / 2,
            )
        )
    linhas.append(
        secao(
            secao=999,
            mesma_secao=True,
            nominais_1t=200,
            lula22_pct=50.0,
            bolso22_pct=50.0,
            lula_pct=20.0,
            flavio_pct=80.0,
        )
    )
    linhas.append(secao(secao=1000, lula_pct=20.0, flavio_pct=80.0))
    j = fc.crit_j(quadro(*linhas))
    assert bool(j["j"].iloc[-2]) is True
    assert bool(j["j"].iloc[-1]) is False  # sem casamento com 2022
    assert int(j["j"].iloc[:200].sum()) <= 2


def test_criterio_k_media_das_outras_secoes() -> None:
    linhas = [
        secao(secao=k + 1, nulos=6 + (k % 3), comparecimento=210, brancos=4)
        for k in range(12)
    ]
    linhas.append(secao(secao=50, nulos=30, comparecimento=210, brancos=4))
    k = fc.crit_k(quadro(*linhas))
    assert k["k"].tolist() == [False] * 12 + [True]
    assert bool(k["k_nulos"].iloc[-1]) and not bool(k["k_brancos"].iloc[-1])


def test_loo_media_e_desvio_das_outras() -> None:
    chave = pd.Series(["z"] * 4)
    p = np.array([1.0, 2.0, 3.0, 10.0])
    media, dp, n = fc._loo(chave, p, np.ones(4, dtype=bool))
    assert media[3] == pytest.approx(2.0)
    assert dp[3] == pytest.approx(1.0)
    assert n[3] == 3


def test_criterio_l_conta_so_criterios_de_secao() -> None:
    df = quadro(
        *(secao(secao=k, local_id="L1") for k in range(4)), secao(local_id="L2")
    )
    marcas = pd.DataFrame({c: [False] * 5 for c in fc.IDS})
    marcas.loc[0:2, "a"] = True
    marcas.loc[3, "g"] = True
    marcas.loc[4, "k"] = True
    lcol = fc.crit_l(df, marcas)
    assert lcol["l"].tolist() == [True, True, True, False, False]


# ---------------------------------------------------------------- pontuação e nível


def test_nivel_alta_so_sem_explicacao_comum_e_enclave_baixa() -> None:
    assert fc.nivel(6, False, False, 5, 3) == "alta"
    assert fc.nivel(6, True, False, 5, 3) == "media"
    assert fc.nivel(3, False, False, 5, 3) == "media"
    assert fc.nivel(3, True, False, 5, 3) == "baixa"
    assert fc.nivel(9, False, True, 5, 3) == "baixa"
    assert fc.nivel(0, False, False, 5, 3) is None


def test_pontuar_pesos_declarados_e_iguais() -> None:
    m = pd.DataFrame({c: [False, False] for c in fc.IDS})
    m["h_tipo"] = ["", "sem_arquivo"]
    m.loc[0, ["a", "b", "g"]] = True
    m.loc[1, "h"] = True
    pont, iguais = fc.pontuar(m)
    assert pont.tolist() == [7, fc.PESOS["h_sem_arquivo"]]
    assert iguais.tolist() == [3, 1]


def test_secoes_sem_arquivo_ficam_em_media() -> None:
    m = fc.finalizar(fc.marcas_sem_arquivo(2), pd.Series([[], []]))
    assert m["nivel"].tolist() == ["media", "media"]
    assert m["criterios"].tolist() == [["h"], ["h"]]


def test_explicacao_comum_e_sem_voto_nao_vira_minuscula() -> None:
    df = quadro(
        secao(tipo_inferido="aldeia ou terra indígena"),
        secao(comparecimento=np.nan),
        secao(comparecimento=30),
        secao(tipo_inferido="zona rural"),
    )
    cods = fb.codigos_explicacao(df)
    assert cods[0] == ["aldeia"]
    assert cods[1] == []
    assert cods[2] == ["minuscula"]
    assert fb.grupo_explicacao(cods[3]) == "exige_explicacao_documental"
    assert fb.grupo_explicacao(cods[0]) == "comum_documentada"
    assert "exige explicação documental" in fb.frase_explicacao([])


# ---------------------------------------------------------------- zona congelada


def test_secoes_congeladas_pega_o_lote_do_ultimo_minuto() -> None:
    horas = ["20:50:00", "20:55:00", "20:59:05", "21:00:09", "21:00:10", "21:05:00"]
    df = quadro(
        *(
            secao(secao=k + 1, _dr=pd.Timestamp(f"2026-10-04 {h}"))
            for k, h in enumerate(horas)
        )
    )
    z = {
        "uf": "SP",
        "mun_tse": "71072",
        "zona": 1,
        "ultima_incompleta_utc": "2026-10-05T00:00:57.000Z",
        "secoes_faltando": 2,
    }
    marca = fb.secoes_congeladas(df, [z])
    assert marca.tolist() == [False, False, False, True, True, False]
    assert z["conferem"] is True
    assert z["secoes"] == [4, 5]
    assert z["ultima_incompleta_brasilia"] == "2026-10-04 21:00:57"


def test_horas_entre_e_conversao() -> None:
    assert fb.horas_entre("2026-10-05T00:00:00.000Z", "2026-10-05T15:00:00Z") == 15
    assert fb.horas_entre("2026-10-05T00:00:00Z", None) is None
    assert fb.utc_para_brasilia("2026-10-05T02:30:00Z") == "2026-10-04 23:30:00"


def test_local_id_e_normalizacao() -> None:
    assert fb.local_id("ma", "07307", 20, 1104.0, 45) == "MA-07307-20-1104"
    assert fb.local_id("ma", "07307", 20, np.nan, 45) == "MA-07307-20-s45"
    assert fb.normalizar("Pau D’Arco do Piauí") == "PAU D'ARCO DO PIAUI"


def test_contexto_por_municipio_filtra_temas() -> None:
    ctx = {
        "itens": [
            {
                "id": "ctx-1",
                "tema": "coercao",
                "municipios": [{"nome": "Rio Largo", "uf": "AL"}],
            },
            {
                "id": "ctx-2",
                "tema": "outro",
                "municipios": [{"nome": "Colina", "uf": "SP"}],
            },
        ]
    }
    m = fb.contexto_por_municipio(ctx)
    assert [c["id"] for c in m[("AL", "RIO LARGO")]] == ["ctx-1"]
    assert ("SP", "COLINA") not in m


# ---------------------------------------------------------------- risco


def test_risco_nulo_nunca_vira_zero() -> None:
    r = ft.risco_local(None, None, None)
    assert r["terra_indigena"] is None
    assert r["homicidios_municipio"]["taxa_100mil"] is None
    assert r["nivel_risco_fiscal"] is None
    assert r["validar"] == "validar com a PM e o TRE local"


def test_nivel_de_risco_por_pontos() -> None:
    lr = {
        "favela": "1",
        "terra_indigena": "0",
        "quilombo": "",
        "prisional": "0",
        "fronteira": "0",
    }
    mr = {
        "homicidios_taxa_100mil": "61,2",
        "homicidios_ano": "2023",
        "homicidios_quintil": "5",
        "crime_organizado": "mapeamento público",
        "crime_organizado_fontes": "Veículo, 2026-10-04, https://exemplo | ctx-033, X",
    }
    r = ft.risco_local(lr, mr, {"sede_km_estrada": 5.0, "sede_km_reta": 4.0})
    assert r["nivel_risco_fiscal"] == "alto"
    assert r["crime_organizado"]["fontes"] == ["Veículo", "ctx-033"]
    assert r["quilombo"] is None
    baixo = ft.risco_local({"favela": "0", "prisional": "0"}, None, None)
    assert baixo["nivel_risco_fiscal"] == "baixo"


def test_limpar_tira_travessao_de_nome_externo() -> None:
    assert ft.limpar("Aeroporto A — B") == "Aeroporto A, B"


def test_haversine_e_mais_proximo() -> None:
    assert ft.haversine_km(0, 0, 0, 1) == pytest.approx(111.19, abs=0.05)
    pontos = [
        {"lat": 0.0, "lon": 1.0, "nome": "x"},
        {"lat": 0.0, "lon": 0.2, "nome": "y"},
    ]
    p, d = ft.mais_proximo(0.0, 0.0, pontos)
    assert p is not None and p["nome"] == "y"
    assert d == pytest.approx(22.24, abs=0.05)


def test_osrm_inativo_nao_consulta_rede(tmp_path: Path) -> None:
    o = ft.Osrm(cache=tmp_path / "c.json", ativo=False)
    assert o.tabela([(0.0, 0.0)], [(0.0, 1.0)]) is None
    assert o.consultas == 0


# ---------------------------------------------------------------- textos e exportáveis


def test_regras_sem_travessao_e_com_rotulos() -> None:
    texto = (
        json.dumps(fr.CRITERIOS, ensure_ascii=False) + fr.AVISO + " ".join(ftx.LIMITES)
    )
    assert "—" not in texto
    assert "não é irregularidade" in fr.AVISO
    assert "não de acusação" in fr.AVISO
    assert "log da urna" in fr.AVISO
    assert [c["id"] for c in fr.CRITERIOS] == list(fc.IDS)


def test_completar_le_o_efeito_de_fechamento() -> None:
    crits = [dict(c) for c in fr.CRITERIOS]
    fech = {
        "voto": {
            "estimadores": [
                {
                    "id": "tarde19_zona",
                    "resultado": {
                        "lula_pp": {"estimativa": 4.861, "ic95": [4.448, 5.293]}
                    },
                }
            ]
        }
    }
    fr.completar(crits, fech)
    d = next(c for c in crits if c["id"] == "d")
    assert "+4,86" in d["explicacao_comum"] and "+5,29" in d["explicacao_comum"]


def _registro(nivel: str, risco: str | None) -> dict[str, object]:
    r = {c[0]: None for c in fx.COLUNAS_SECAO}
    base = {
        "nivel": nivel,
        "pontuacao": 6,
        "uf": "SP",
        "municipio": "SÃO PAULO",
        "mun_tse": "71072",
        "ibge": "3550308",
        "zona": 1,
        "secao": 10,
        "local": "ESCOLA",
        "endereco": "RUA A",
        "bairro": "CENTRO",
        "cep": "01000000",
        "lat": -23.5,
        "lon": -46.6,
        "lula_pct": 12.5,
        "criterios": ["a", "b"],
        "nivel_iguais": "media",
        "enclave_2022": False,
        "sem_boletim": False,
        "contexto": [],
        "risco": {
            "nivel_risco_fiscal": risco,
            "motivos_risco": [],
            "validar": ft.VALIDAR,
        },
        "link_mapa": {
            "osm": "https://www.openstreetmap.org/?mlat=-23.5",
            "google": "https://www.google.com/maps?q=-23.5,-46.6",
        },
    }
    r.update(base)
    for k in ("flavio_2022_pct", "lula_2022_pct", "excesso_zona_lula_pp"):
        r.setdefault(k, None)
    return r


def test_csv_com_virgula_decimal(tmp_path: Path) -> None:
    caminho = tmp_path / "s.csv"
    regs = [_registro("alta", "alto")]
    for k in ("lula", "flavio", "flavio_2022_pct", "lula_2022_pct"):
        regs[0][k] = None
    info = fx.escrever_csv(caminho, regs, fx.COLUNAS_SECAO)
    assert info["linhas"] == 1
    with caminho.open(encoding="utf-8-sig") as f:
        linhas = list(csv.reader(f, delimiter=";"))
    cab, val = linhas
    assert val[cab.index("lula_pct")] == "12,5"
    assert val[cab.index("criterios")] == "a,b"


def test_excel_abas_cores_e_linhas(tmp_path: Path) -> None:
    regs = [_registro("alta", "alto"), _registro("baixa", None)]
    for r in regs:
        for k in ("lula", "flavio"):
            r[k] = None
    dados = {
        "titulo": "t",
        "gerado_em": "2026-10-06T00:00:00Z",
        "meta": {
            "data_eleicao": "2026-10-04",
            "segundo_turno": "2026-10-25",
            "cortes_nivel": fr.cortes_json(),
            "pesos": fr.pesos_json(),
            "fontes": [{"chave": "x", "caminho": "y", "descricao": "z", "sha256": "0"}],
            "fontes_risco": [],
            "risco_regra": ft.regra(),
        },
        "secoes": regs,
        "por_local": [],
        "por_uf": [],
        "criterios": [
            c
            | {
                "peso": 1,
                "secoes": 0,
                "secoes_por_nivel": {"alta": 0, "media": 0, "baixa": 0},
            }
            for c in fr.CRITERIOS
        ],
    }
    caminho = tmp_path / "f.xlsx"
    fx.escrever_excel(caminho, dados, [])
    wb = load_workbook(caminho)
    assert wb.sheetnames == [
        "Leia-me",
        "Seções",
        "Locais",
        "Municípios",
        "UFs",
        "Critérios",
        "Riscos",
    ]
    ws = wb["Seções"]
    assert ws.max_row == 1 + len(regs)
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert ws["A2"].fill.start_color.rgb.endswith(fx.COR_NIVEL["alta"])
    col_link = [c.value for c in ws[1]].index("link_openstreetmap") + 1
    assert ws.cell(2, col_link).hyperlink is not None


# ---------------------------------------------------------------- arquivos publicados


@pytest.mark.skipif(not JSON.exists(), reason="fiscais.json ainda não gerado")
def test_json_publicado_segue_o_contrato() -> None:
    d = json.loads(JSON.read_text(encoding="utf-8"))
    assert d["versao_contrato"] == "1.1"
    assert "—" not in JSON.read_text(encoding="utf-8")
    assert set(d["rotulos"]) == {"atipico", "prioridade", "resolve"}
    niveis = [s["nivel"] for s in d["secoes"]]
    ordem = [fc.ORDEM_NIVEL[n] for n in niveis]
    assert ordem == sorted(ordem)
    for s in d["secoes"]:
        assert s["criterios"], s
        if s["nivel"] == "alta":
            assert s["grupo_explicacao"] == "exige_explicacao_documental"
        if s["enclave_2022"]:
            assert s["nivel"] == "baixa"
    assert d["resumo"]["secoes_sinalizadas"] == len(d["secoes"])
    assert sum(v["secoes"] for v in d["resumo"]["por_nivel"].values()) == len(
        d["secoes"]
    )
    assert d["meta"]["mistura"]["conferencia"]["iguais"] == 50
    assert len(d["sem_arquivo"]) == d["meta"]["universo"]["secoes_sem_arquivo"]
    assert all(z["conferem"] for z in d["zonas_congeladas"])
    assert len(d["mapa"]["pontos"]) + d["mapa"]["n_sem_coordenada"] == len(
        d["por_local"]
    )
    exp = {e["formato"] + e["caminho"][-8:]: e for e in d["meta"]["exportaveis"]}
    assert len(exp) == 3


@pytest.mark.skipif(
    not (JSON.exists() and XLSX.exists()), reason="exportáveis ausentes"
)
def test_excel_publicado_tem_as_linhas_do_json() -> None:
    d = json.loads(JSON.read_text(encoding="utf-8"))
    wb = load_workbook(XLSX, read_only=True)
    ws = wb["Seções"]
    assert ws.max_row - 1 == len(d["secoes"])
    assert wb.sheetnames == [
        "Leia-me",
        "Seções",
        "Locais",
        "Municípios",
        "UFs",
        "Critérios",
        "Riscos",
    ]
    with CSV.open(encoding="utf-8-sig") as f:
        assert sum(1 for _ in f) - 1 == len(d["secoes"])
    xl = next(e for e in d["meta"]["exportaveis"] if e["formato"] == "xlsx")
    assert xl["bytes"] == XLSX.stat().st_size
