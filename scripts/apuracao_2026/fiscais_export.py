"""Exportáveis do capítulo 13: CSV por seção, CSV por local e Excel com sete abas.

CSV com `;` como separador, vírgula decimal e UTF-8 com BOM, para abrir direto no
Excel em português. O Excel (openpyxl) tem cabeçalho congelado e em negrito, filtro
automático em cada aba, larguras ajustadas, formatos numéricos, links clicáveis e
cor de fundo por nível de prioridade e por nível de risco fiscal.
"""

from __future__ import annotations

import csv
import hashlib
from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .fiscais_regras import AVISO, CRITERIOS, POR_ID, ROTULOS

NOME_NIVEL = {"alta": "alta", "media": "média", "baixa": "baixa"}
NOME_RISCO = {"alto": "alto", "medio": "médio", "baixo": "baixo"}
COR_NIVEL = {"alta": "FFC7CE", "media": "FFEB9C", "baixa": "D9D9D9"}
COR_RISCO = {"alto": "F4B084", "medio": "FFF2CC", "baixo": "E2EFDA"}


def _sim(v: Any) -> str | None:
    if v is None:
        return None
    return "sim" if v else "não"


def _risco(s: Mapping[str, Any], *caminho: str) -> Any:
    r: Any = s.get("risco")
    for c in caminho:
        if not isinstance(r, Mapping):
            return None
        r = r.get(c)
    return r


# (cabeçalho, extrator, formato numérico ou None, largura)
Coluna = tuple[str, Callable[[Mapping[str, Any]], Any], str | None, int]

COLUNAS_SECAO: list[Coluna] = [
    ("nivel", lambda s: NOME_NIVEL.get(s["nivel"], s["nivel"]), None, 8),
    ("pontuacao", lambda s: s["pontuacao"], "0", 10),
    ("uf", lambda s: s["uf"], None, 5),
    ("municipio", lambda s: s["municipio"], None, 24),
    ("mun_tse", lambda s: s["mun_tse"], None, 8),
    ("ibge", lambda s: s["ibge"], None, 9),
    ("zona", lambda s: s["zona"], "0", 6),
    ("secao", lambda s: s["secao"], "0", 7),
    ("local", lambda s: s["local"], None, 36),
    ("endereco", lambda s: s["endereco"], None, 34),
    ("bairro", lambda s: s["bairro"], None, 20),
    ("cep", lambda s: s["cep"], None, 10),
    ("lat", lambda s: s["lat"], "0.00000", 10),
    ("lon", lambda s: s["lon"], "0.00000", 10),
    ("tipo_local_inferido", lambda s: s["tipo_local_inferido"], None, 18),
    ("aptos", lambda s: s["aptos"], "#,##0", 8),
    ("votantes", lambda s: s["votantes"], "#,##0", 9),
    ("validos", lambda s: s["validos"], "#,##0", 8),
    ("lula_votos", lambda s: s["lula"], "#,##0", 9),
    ("flavio_votos", lambda s: s["flavio"], "#,##0", 9),
    ("lula_pct", lambda s: s["lula_pct"], "0.0", 8),
    ("flavio_pct", lambda s: s["flavio_pct"], "0.0", 8),
    ("zona_lula_pct", lambda s: s["zona_lula_pct"], "0.0", 9),
    ("zona_flavio_pct", lambda s: s["zona_flavio_pct"], "0.0", 9),
    ("excesso_zona_lula_pp", lambda s: s["excesso_zona_lula_pp"], "+0.0;-0.0;0.0", 10),
    (
        "excesso_zona_flavio_pp",
        lambda s: s["excesso_zona_flavio_pp"],
        "+0.0;-0.0;0.0",
        10,
    ),
    ("lula_2022_pct", lambda s: s["lula_2022_pct"], "0.0", 9),
    ("bolsonaro_2022_pct", lambda s: s["flavio_2022_pct"], "0.0", 9),
    ("modelo_urna", lambda s: s["modelo_urna"], None, 9),
    ("tipo_urna", lambda s: s["tipo_urna"], "0", 6),
    ("tipo_arquivo", lambda s: s["tipo_arquivo"], "0", 6),
    ("n_cargas", lambda s: s["n_cargas"], "0", 6),
    ("encerramento_brasilia", lambda s: s["encerramento_brasilia"], None, 19),
    ("recebido_tse", lambda s: s["recebido_tse"], None, 19),
    ("criterios", lambda s: ",".join(s["criterios"]), None, 9),
    (
        "criterios_por_extenso",
        lambda s: "; ".join(f"{c}: {POR_ID[c]['nome']}" for c in s["criterios"]),
        None,
        48,
    ),
    ("nivel_pesos_iguais", lambda s: NOME_NIVEL.get(s["nivel_iguais"]), None, 9),
    ("grupo_explicacao", lambda s: s["grupo_explicacao"], None, 22),
    ("explicacao_provavel", lambda s: s["explicacao_provavel"], None, 40),
    ("enclave_2022", lambda s: _sim(s["enclave_2022"]), None, 7),
    ("sem_boletim", lambda s: _sim(s["sem_boletim"]), None, 7),
    ("o_que_conferir", lambda s: s["o_que_conferir"], None, 60),
    ("contexto_imprensa", lambda s: ",".join(s["contexto"]), None, 14),
    ("risco_nivel", lambda s: NOME_RISCO.get(_risco(s, "nivel_risco_fiscal")), None, 8),
    ("risco_motivos", lambda s: "; ".join(_risco(s, "motivos_risco") or []), None, 40),
    ("risco_validar", lambda s: _risco(s, "validar"), None, 22),
    ("rural_urbano", lambda s: _risco(s, "rural_urbano"), None, 8),
    ("terra_indigena", lambda s: _sim(_risco(s, "terra_indigena")), None, 7),
    ("quilombo", lambda s: _sim(_risco(s, "quilombo")), None, 7),
    ("favela_comunidade", lambda s: _sim(_risco(s, "favela_comunidade")), None, 8),
    (
        "unidade_prisional",
        lambda s: _sim(_risco(s, "unidade_prisional_ou_socioeducativa")),
        None,
        8,
    ),
    (
        "homicidios_taxa_100mil",
        lambda s: _risco(s, "homicidios_municipio", "taxa_100mil"),
        "0.0",
        9,
    ),
    ("homicidios_ano", lambda s: _risco(s, "homicidios_municipio", "ano"), "0", 6),
    (
        "homicidios_quintil",
        lambda s: _risco(s, "homicidios_municipio", "quintil"),
        "0",
        6,
    ),
    ("crime_organizado", lambda s: _risco(s, "crime_organizado", "status"), None, 16),
    (
        "fronteira",
        lambda s: _sim(_risco(s, "fronteira_ou_garimpo", "fronteira")),
        None,
        7,
    ),
    ("garimpo", lambda s: _sim(_risco(s, "fronteira_ou_garimpo", "garimpo")), None, 7),
    (
        "sede_km_estrada",
        lambda s: _risco(s, "acesso", "sede_km_estrada"),
        "#,##0.0",
        9,
    ),
    ("sede_km_reta", lambda s: _risco(s, "acesso", "sede_km_reta"), "#,##0.0", 9),
    ("aeroporto", lambda s: _risco(s, "acesso", "aeroporto"), None, 22),
    (
        "aeroporto_km_estrada",
        lambda s: _risco(s, "acesso", "aeroporto_km_estrada"),
        "#,##0.0",
        9,
    ),
    (
        "aeroporto_km_reta",
        lambda s: _risco(s, "acesso", "aeroporto_km_reta"),
        "#,##0.0",
        9,
    ),
    ("link_openstreetmap", lambda s: (s["link_mapa"] or {}).get("osm"), "link", 14),
    ("link_google_maps", lambda s: (s["link_mapa"] or {}).get("google"), "link", 14),
]

COLUNAS_LOCAL: list[Coluna] = [
    ("nivel", lambda r: NOME_NIVEL.get(r["nivel"], r["nivel"]), None, 8),
    ("pontuacao_soma", lambda r: r["pontuacao_soma"], "0", 10),
    ("pontuacao_max", lambda r: r["pontuacao_max"], "0", 8),
    ("secoes_sinalizadas", lambda r: r["secoes"], "0", 9),
    ("secoes_no_local", lambda r: r["secoes_local"], "0", 9),
    ("lista_secoes", lambda r: ",".join(str(x) for x in r["lista_secoes"]), None, 18),
    ("alta", lambda r: r["alta"], "0", 5),
    ("media", lambda r: r["media"], "0", 6),
    ("baixa", lambda r: r["baixa"], "0", 6),
    ("uf", lambda r: r["uf"], None, 5),
    ("municipio", lambda r: r["municipio"], None, 24),
    ("mun_tse", lambda r: r["mun_tse"], None, 8),
    ("zona", lambda r: r["zona"], "0", 6),
    ("local_nr", lambda r: r["local_nr"], "0", 7),
    ("local", lambda r: r["local"], None, 36),
    ("endereco", lambda r: r["endereco"], None, 34),
    ("bairro", lambda r: r["bairro"], None, 20),
    ("cep", lambda r: r["cep"], None, 10),
    ("lat", lambda r: r["lat"], "0.00000", 10),
    ("lon", lambda r: r["lon"], "0.00000", 10),
    ("aptos_local", lambda r: r["aptos_local"], "#,##0", 9),
    ("aptos_secoes_sinalizadas", lambda r: r["aptos_secoes"], "#,##0", 9),
    ("lula_pct_local", lambda r: r["lula_pct"], "0.0", 8),
    ("flavio_pct_local", lambda r: r["flavio_pct"], "0.0", 8),
    (
        "criterios",
        lambda r: "; ".join(f"{k}: {v}" for k, v in r["criterios"].items()),
        None,
        24,
    ),
    ("contexto_imprensa", lambda r: ",".join(r["contexto"]), None, 14),
    ("risco_nivel", lambda r: NOME_RISCO.get(_risco(r, "nivel_risco_fiscal")), None, 8),
    ("risco_motivos", lambda r: "; ".join(_risco(r, "motivos_risco") or []), None, 40),
    ("link_openstreetmap", lambda r: (r["link_mapa"] or {}).get("osm"), "link", 14),
    ("link_google_maps", lambda r: (r["link_mapa"] or {}).get("google"), "link", 14),
]

COLUNAS_MUNICIPIO: list[Coluna] = [
    ("uf", lambda r: r["uf"], None, 5),
    ("municipio", lambda r: r["municipio"], None, 26),
    ("mun_tse", lambda r: r["mun_tse"], None, 8),
    ("ibge", lambda r: r["ibge"], None, 9),
    ("secoes_sinalizadas", lambda r: r["secoes"], "#,##0", 10),
    ("secoes_universo", lambda r: r["secoes_universo"], "#,##0", 10),
    ("alta", lambda r: r["alta"], "0", 6),
    ("media", lambda r: r["media"], "0", 6),
    ("baixa", lambda r: r["baixa"], "#,##0", 6),
    ("locais", lambda r: r["locais"], "#,##0", 7),
    ("aptos_secoes_sinalizadas", lambda r: r["aptos_secoes"], "#,##0", 10),
    ("pontuacao_soma", lambda r: r["pontuacao_soma"], "#,##0", 10),
    ("lula_pct_municipio", lambda r: r["lula_pct"], "0.0", 8),
    ("flavio_pct_municipio", lambda r: r["flavio_pct"], "0.0", 8),
    ("margem_flavio_pp", lambda r: r["margem_flavio_pp"], "+0.0;-0.0;0.0", 9),
    (
        "criterios",
        lambda r: "; ".join(f"{k}: {v}" for k, v in r["criterios"].items()),
        None,
        30,
    ),
    (
        "contexto_imprensa",
        lambda r: ", ".join(f"{c['id']} ({c['tema']})" for c in r["contexto"]),
        None,
        30,
    ),
]

COLUNAS_UF: list[Coluna] = [
    ("uf", lambda r: r["uf"], None, 5),
    ("regiao", lambda r: r["regiao"], None, 12),
    ("secoes_sinalizadas", lambda r: r["secoes"], "#,##0", 10),
    ("secoes_universo", lambda r: r["secoes_universo"], "#,##0", 10),
    ("alta", lambda r: r["alta"], "0", 6),
    ("media", lambda r: r["media"], "#,##0", 6),
    ("baixa", lambda r: r["baixa"], "#,##0", 7),
    ("locais", lambda r: r["locais"], "#,##0", 7),
    ("locais_alta", lambda r: r["locais_alta"], "0", 7),
    ("aptos_secoes_sinalizadas", lambda r: r["aptos_secoes"], "#,##0", 11),
    ("pontuacao_soma", lambda r: r["pontuacao_soma"], "#,##0", 10),
    (
        "criterios",
        lambda r: "; ".join(f"{k}: {v}" for k, v in r["criterios"].items()),
        None,
        40,
    ),
]


# ---------------------------------------------------------------- CSV


def _csv_valor(v: Any) -> Any:
    if isinstance(v, float):
        return f"{v}".replace(".", ",")
    if isinstance(v, bool):
        return "sim" if v else "não"
    return "" if v is None else v


def escrever_csv(
    caminho: Path, linhas: Sequence[Mapping[str, Any]], colunas: Sequence[Coluna]
) -> dict[str, Any]:
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\n")
        w.writerow([c[0] for c in colunas])
        for r in linhas:
            w.writerow([_csv_valor(c[1](r)) for c in colunas])
    return {
        "linhas": len(linhas),
        "colunas": [c[0] for c in colunas],
        "separador": ";",
        "decimal": ",",
        "codificacao": "utf-8-sig",
    }


def assinatura(caminho: Path) -> dict[str, Any]:
    h = hashlib.sha256(caminho.read_bytes()).hexdigest()
    return {"bytes": caminho.stat().st_size, "sha256": h}


# ---------------------------------------------------------------- Excel

NEGRITO = Font(bold=True)
LINK = Font(color="0563C1", underline="single")


def _aba(
    wb: Workbook,
    titulo: str,
    linhas: Sequence[Mapping[str, Any]],
    colunas: Sequence[Coluna],
    cor_linha: Callable[[Mapping[str, Any]], tuple[int, str] | None] | None = None,
    extras: Sequence[Callable[[Mapping[str, Any]], tuple[int, str] | None]] = (),
) -> None:
    ws = wb.create_sheet(titulo)
    ws.append([c[0] for c in colunas])
    for cel in ws[1]:
        cel.font = NEGRITO
        cel.alignment = Alignment(vertical="top", wrap_text=True)
    fills: dict[str, PatternFill] = {}

    def fill(cor: str) -> PatternFill:
        if cor not in fills:
            fills[cor] = PatternFill("solid", start_color=cor, end_color=cor)
        return fills[cor]

    links = [i for i, c in enumerate(colunas, start=1) if c[2] == "link"]
    for r in linhas:
        ws.append([c[1](r) for c in colunas])
        n = ws.max_row
        for i, c in enumerate(colunas, start=1):
            fmt = c[2]
            if fmt and fmt != "link":
                ws.cell(n, i).number_format = fmt
        for i in links:
            cel = ws.cell(n, i)
            if cel.value:
                cel.hyperlink = cel.value
                cel.font = LINK
        for f in (cor_linha, *extras):
            if f is None:
                continue
            res = f(r)
            if res:
                col, cor = res
                ws.cell(n, col).fill = fill(cor)
    for i, c in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = c[3]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(ws.max_row, 1)}"


def _texto(
    wb: Workbook, titulo: str, linhas: Sequence[Sequence[Any]], larguras: Sequence[int]
) -> None:
    ws = wb.create_sheet(titulo)
    for r in linhas:
        ws.append(list(r))
    for cel in ws[1]:
        cel.font = NEGRITO
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for linha in ws.iter_rows(min_row=1):
        for cel in linha:
            cel.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(larguras))}{max(ws.max_row, 1)}"


def _peso(p: Any) -> str:
    if isinstance(p, Mapping):
        return " ou ".join(f"{v} ({k.replace('_', ' ')})" for k, v in p.items())
    return str(p)


def _col(colunas: Sequence[Coluna], nome: str) -> int:
    return next(i for i, c in enumerate(colunas, start=1) if c[0] == nome)


def escrever_excel(
    caminho: Path,
    dados: Mapping[str, Any],
    municipios: Sequence[Mapping[str, Any]],
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    meta = dados["meta"]
    leia = [
        ["Campo", "Conteúdo"],
        ["O que é", dados["titulo"]],
        ["Rótulo 1", ROTULOS["atipico"]],
        ["Rótulo 2", ROTULOS["prioridade"]],
        ["Rótulo 3", ROTULOS["resolve"]],
        ["Gerado em (UTC)", dados["gerado_em"]],
        [
            "Eleição",
            f"1º turno em {meta['data_eleicao']}; 2º turno em {meta['segundo_turno']}",
        ],
        [
            "Como usar",
            "Aba Seções: uma linha por seção sinalizada, da prioridade alta para a baixa. "
            "Aba Locais: o mesmo agrupado por local de votação (um fiscal pode cobrir "
            "várias seções do mesmo local). Filtre por UF e município, abra o link do "
            "mapa e leve ao local a coluna o_que_conferir.",
        ],
        ["Níveis", meta["cortes_nivel"]["regra"]],
        [
            "Cores",
            "nível: alta vermelho claro, média amarelo claro, baixa cinza; risco fiscal: "
            "alto laranja, médio amarelo, baixo verde claro",
        ],
        ["Pesos", "; ".join(f"{k}: {v}" for k, v in meta["pesos"].items())],
    ]
    for c in CRITERIOS:
        leia.append([f"Critério {c['id']}", f"{c['nome']}: {c['regra']}"])
    for f in meta["fontes"]:
        leia.append(
            [
                f"Fonte: {f['chave']}",
                f"{f['caminho']}; {f['descricao']}; SHA-256 {f.get('sha256')}",
            ]
        )
    leia.append(["Aviso", AVISO])
    _texto(wb, "Leia-me", leia, [22, 120])

    ic, ir = _col(COLUNAS_SECAO, "nivel"), _col(COLUNAS_SECAO, "risco_nivel")
    _aba(
        wb,
        "Seções",
        dados["secoes"],
        COLUNAS_SECAO,
        lambda s: (ic, COR_NIVEL[s["nivel"]]) if s["nivel"] in COR_NIVEL else None,
        (
            lambda s: (
                (ir, COR_RISCO[k])
                if (k := _risco(s, "nivel_risco_fiscal")) in COR_RISCO
                else None
            ),
        ),
    )
    jl, jr = _col(COLUNAS_LOCAL, "nivel"), _col(COLUNAS_LOCAL, "risco_nivel")
    _aba(
        wb,
        "Locais",
        dados["por_local"],
        COLUNAS_LOCAL,
        lambda r: (jl, COR_NIVEL[r["nivel"]]) if r["nivel"] in COR_NIVEL else None,
        (
            lambda r: (
                (jr, COR_RISCO[k])
                if (k := _risco(r, "nivel_risco_fiscal")) in COR_RISCO
                else None
            ),
        ),
    )
    _aba(wb, "Municípios", municipios, COLUNAS_MUNICIPIO)
    _aba(wb, "UFs", dados["por_uf"], COLUNAS_UF)
    crit = [
        [
            "id",
            "nome",
            "o que mede",
            "regra",
            "limiar",
            "peso",
            "seções",
            "alta",
            "média",
            "baixa",
            "explicação comum",
            "o que o fiscal confere",
            "fonte",
        ]
    ]
    for c in dados["criterios"]:
        pn = c["secoes_por_nivel"]
        crit.append(
            [
                c["id"],
                c["nome"],
                c["mede"],
                c["regra"],
                c["limiar"],
                _peso(c["peso"]),
                c["secoes"],
                pn["alta"],
                pn["media"],
                pn["baixa"],
                c["explicacao_comum"],
                c["o_que_conferir"],
                c["fonte"],
            ]
        )
    _texto(wb, "Critérios", crit, [4, 26, 30, 60, 22, 8, 8, 6, 6, 7, 34, 40, 30])
    regra = meta.get("risco_regra") or {}
    riscos = [
        [
            "camada",
            "regra",
            "fonte",
            "url",
            "baixado em",
            "SHA-256",
            "status",
            "motivo",
            "cobertura (seções)",
        ]
    ]
    for f in meta.get("fontes_risco", []):
        riscos.append(
            [
                f.get("camada"),
                f.get("regra"),
                f.get("nome"),
                f.get("url"),
                f.get("baixado_em"),
                f.get("sha256"),
                f.get("status"),
                f.get("motivo"),
                f.get("cobertura_secoes"),
            ]
        )
    riscos.append(
        [
            "nível de risco fiscal",
            regra.get("texto"),
            "regra declarada (juízo editorial)",
            None,
            None,
            None,
            None,
            regra.get("frase"),
            None,
        ]
    )
    _texto(wb, "Riscos", riscos, [22, 50, 30, 40, 18, 20, 8, 30, 10])
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)
